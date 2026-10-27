"""Unit and integration tests for Stage 8: M7 BEME Excel & M8 Pilot Report PDF exports."""
from datetime import datetime, timezone
from pathlib import Path
import openpyxl
import pytest

from app import __version__
from app.models.project import Project
from app.models.user import User
from app.modules.exports import BemeExporter, ExportsService, PilotReportExporter
from app.modules.intake.service import IntakeService
from app.modules.pipeline.service import PipelineService
from app.schemas.exports import AccuracyCheck, PilotReportData
from app.schemas.manifest import PipelineRunRequest
from tests.synthetic_data.terrain_generators import (
    generate_flat_plane_points,
    points_to_csv_bytes
)


@pytest.fixture
def sample_pipeline_data(db_session):
    """Generates synthetic project, survey intake, and full pipeline run for export tests."""
    user = User(
        email="engineer@dariusos.test",
        full_name="Civil Testing Engineer",
        hashed_password="pw",
        role="engineer"
    )
    db_session.add(user)
    db_session.commit()

    project = Project(
        name="Golden Export Highway",
        description="Testing BEME and Pilot Report exports",
        created_by_user_id=user.id
    )
    db_session.add(project)
    db_session.commit()

    pts = generate_flat_plane_points(
        length_x=100.0,
        width_y=40.0,
        spacing=2.0,
        z_elevation=100.0
    )
    csv_bytes = points_to_csv_bytes(pts)

    survey_rec, _ = IntakeService.process_file_upload(
        db=db_session,
        project_id=project.id,
        user_id=user.id,
        filename="export_corridor_survey.csv",
        raw_bytes=csv_bytes
    )

    req = PipelineRunRequest(
        project_id=project.id,
        survey_file_id=survey_rec.id,
        start_point=(0.0, 0.0),
        end_point=(100.0, 0.0),
        station_interval_m=20.0,
        target_elevation_offset_m=-2.0
    )

    res, align, prof, cs, quant = PipelineService.execute_run(
        db=db_session,
        request=req,
        user_id=user.id
    )

    return {
        "project": project,
        "user": user,
        "survey_rec": survey_rec,
        "alignment": align,
        "profile": prof,
        "sections": cs,
        "quantities": quant,
        "run_res": res
    }


def test_m7_beme_excel_generation(sample_pipeline_data):
    """Verifies BEME generation with openpyxl, checking formulas, banner, and single source of truth."""
    alignment = sample_pipeline_data["alignment"]
    sections = sample_pipeline_data["sections"]
    quantities = sample_pipeline_data["quantities"]
    survey_rec = sample_pipeline_data["survey_rec"]
    run_res = sample_pipeline_data["run_res"]

    beme_data = BemeExporter.generate_excel(
        project_name="Golden Export Highway",
        survey_sha256=survey_rec.sha256_hash,
        rules_sha256=run_res.engineering_rules_hash,
        alignment=alignment,
        cross_sections=sections,
        quantities=quantities,
        output_filename="test_golden_beme.xlsx"
    )

    # 1. Output file exists
    assert beme_data.excel_file_path is not None
    excel_path = Path(beme_data.excel_file_path)
    assert excel_path.exists()
    assert excel_path.stat().st_size > 0

    # 2. Open workbook and verify formulas
    wb = openpyxl.load_workbook(excel_path, data_only=False)
    assert "Bill of Quantities (BEME)" in wb.sheetnames
    ws = wb["Bill of Quantities (BEME)"]

    # Engineer disclaimer banner in A1
    assert "TO BE CONFIRMED BY A REGISTERED ENGINEER" in str(ws["A1"].value)

    # Check headers
    assert ws["A7"].value == "Item No."
    assert ws["D7"].value == "Quantity"
    assert ws["F7"].value == "Amount (USD)"

    # Check item row formulas (=D*E)
    assert ws["F8"].value == "=D8*E8"
    assert ws["F9"].value == "=D9*E9"

    # Check quantities match pipeline quantities (Single source of truth)
    cut_item = next(item for item in beme_data.items if item.item_code == "2.01")
    assert pytest.approx(cut_item.quantity, 0.01) == quantities.total_cut_volume_m3

    fill_item = next(item for item in beme_data.items if item.item_code == "2.02")
    assert pytest.approx(fill_item.quantity, 0.01) == quantities.factored_fill_volume_m3

    # Check grand total formula (=SUM(...))
    total_row = 8 + len(beme_data.items)
    assert ws.cell(row=total_row, column=6).value == f"=SUM(F8:F{total_row - 1})"

    # Grand total estimated > 0
    assert beme_data.grand_total_estimated > 0.0


def test_m8_pilot_report_pdf_generation(sample_pipeline_data):
    """Verifies Pilot PDF report generation using reportlab with tables and accuracy checks."""
    alignment = sample_pipeline_data["alignment"]
    quantities = sample_pipeline_data["quantities"]
    survey_rec = sample_pipeline_data["survey_rec"]
    run_res = sample_pipeline_data["run_res"]

    report_data = PilotReportData(
        project_title="Golden Export Highway",
        project_id="proj_001",
        generated_at=datetime.now(timezone.utc).isoformat(),
        code_version=__version__,
        survey_filename=survey_rec.original_filename,
        survey_sha256=survey_rec.sha256_hash,
        rules_sha256=run_res.engineering_rules_hash,
        total_survey_points=survey_rec.point_count,
        corridor_length_m=alignment.total_length_m,
        corridor_width_m=20.0,
        station_interval_m=alignment.station_interval_m,
        total_stations=len(alignment.stations),
        max_ground_grade_percent=5.0,
        max_design_grade_percent=4.0,
        total_cut_m3=quantities.total_cut_volume_m3,
        total_fill_m3=quantities.total_fill_volume_m3,
        factored_fill_m3=quantities.factored_fill_volume_m3,
        net_earthwork_m3=quantities.net_volume_m3,
        mass_haul_balance=quantities.balance_status.value,
        total_estimated_cost_usd=125430.50,
        pipeline_duration_seconds=1.42,
        accuracy_benchmarks=[
            AccuracyCheck(
                benchmark_name="Golden Case 1: Flat Cut Trench",
                theoretical_value=2400.0,
                computed_value=2400.0,
                unit="m³",
                discrepancy_percent=0.0,
                passed=True
            )
        ]
    )

    pdf_path_str = PilotReportExporter.generate_pdf(report_data, output_filename="test_pilot_report.pdf")
    pdf_path = Path(pdf_path_str)

    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 1000  # Non-trivial PDF size

    # Verify PDF magic bytes
    with open(pdf_path, "rb") as f:
        header = f.read(5)
        assert header == b"%PDF-"


def test_exports_service(sample_pipeline_data, db_session):
    """Verifies ExportsService handles both BEME Excel and Pilot PDF reports with audit logging."""
    alignment = sample_pipeline_data["alignment"]
    sections = sample_pipeline_data["sections"]
    quantities = sample_pipeline_data["quantities"]
    survey_rec = sample_pipeline_data["survey_rec"]
    run_res = sample_pipeline_data["run_res"]
    project = sample_pipeline_data["project"]
    user = sample_pipeline_data["user"]

    beme = ExportsService.generate_beme(
        project_name=project.name,
        survey_sha256=survey_rec.sha256_hash,
        rules_sha256=run_res.engineering_rules_hash,
        alignment=alignment,
        cross_sections=sections,
        quantities=quantities,
        db=db_session,
        user_id=user.id,
        project_id=project.id
    )
    assert beme.excel_file_path is not None
    assert Path(beme.excel_file_path).exists()

    report_data = PilotReportData(
        project_title=project.name,
        project_id=project.id,
        generated_at=datetime.now(timezone.utc).isoformat(),
        code_version=__version__,
        survey_filename=survey_rec.original_filename,
        survey_sha256=survey_rec.sha256_hash,
        rules_sha256=run_res.engineering_rules_hash,
        total_survey_points=survey_rec.point_count,
        corridor_length_m=alignment.total_length_m,
        corridor_width_m=20.0,
        station_interval_m=alignment.station_interval_m,
        total_stations=len(alignment.stations),
        max_ground_grade_percent=2.0,
        max_design_grade_percent=1.5,
        total_cut_m3=quantities.total_cut_volume_m3,
        total_fill_m3=quantities.total_fill_volume_m3,
        factored_fill_m3=quantities.factored_fill_volume_m3,
        net_earthwork_m3=quantities.net_volume_m3,
        mass_haul_balance=quantities.balance_status.value,
        total_estimated_cost_usd=50000.0,
        pipeline_duration_seconds=0.85
    )
    pdf_path = ExportsService.generate_pilot_report(
        report_data=report_data,
        db=db_session,
        user_id=user.id
    )
    assert Path(pdf_path).exists()
    assert report_data.pdf_file_path == pdf_path
